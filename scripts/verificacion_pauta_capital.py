#!/usr/bin/env python3
"""
Archivo de verificación de la pauta de Capital: una fila por cada fila real del crudo, con
su tipo de geografía, su destino, si entra al tablero (y el motivo si no), cuántas entradas
genera, su estado, y si es activa / activa dirigida en cada fecha de control. Sirve para
cotejar cualquier conteo del tablero fila por fila (mismo papel que el archivo de apoyo del
tablero del interior).

Uso:
    python3 scripts/verificacion_pauta_capital.py --crudo "<ruta>/Pauta RRSS_GLOBAL.xlsx" \
        --salida Verificacion_pauta_capital.xlsx [--fechas 2026-10-14 2026-10-15]

Usa las mismas funciones que el generador (no hay una segunda copia de las reglas).
"""
import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import generar_pauta_capital as g  # noqa: E402


def tipo_geografia(loc):
    if loc is None or (isinstance(loc, float) and pd.isna(loc)):
        return "Vacío"
    loc = str(loc).strip()
    if g.CODIGO_CIRCUITO_RE.match(loc):
        return "Circuito de Capital (código)"
    if loc == g.TODAS:
        return "TODAS LAS LOCALIDADES"
    if g.normalizar(loc.split(",")[0]) == "CORDOBA CAPITAL":
        return "CÓRDOBA CAPITAL (la audiencia decide)"
    return "Otra localidad"


def _a_fecha(v):
    """Fecha real de Excel (date) o None: así las fórmulas del manual pueden compararla."""
    ts = g._a_timestamp(v)
    return None if ts is None else ts.date()


def clave_pieza(p):
    return json.dumps([p["titulo_pieza"], p["campania"], p["plataforma"], p["inicio"], p["fin"], p["objetivo_campana"],
                       p["tipo_pauta"], p["estado"], p["material"]["link"], p["metricas_ejecutadas"]], ensure_ascii=False)


def activa(p, fecha):
    """Mismo predicado que lib/estadoPauta.js: estado Vigente e INICIO <= fecha <= FIN."""
    return p["estado"] == "Vigente" and bool(p["inicio"]) and bool(p["fin"]) and p["inicio"] <= fecha <= p["fin"]


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--crudo", default=str(g.CRUDO_DEFECTO), help="por defecto Contexto/Pauta RRSS_GLOBAL.xlsx")
    ap.add_argument("--salida", default=str(g.ROOT / "Contexto" / "Verificacion_pauta_capital.xlsx"), help="por defecto Contexto/Verificacion_pauta_capital.xlsx (no se versiona)")
    ap.add_argument("--fechas", nargs="+", default=["2026-10-14", "2026-10-15"])
    args = ap.parse_args(argv)

    ref = g.cargar_referencias()
    df_crudo = pd.read_excel(args.crudo, sheet_name="GLOBAL_ImportRange")
    salida, stats, _ = g.generar(df_crudo, ref)
    por_fila = defaultdict(list)
    for _, _, p in g.entradas_de(salida):
        por_fila[p["fila_origen"]].append(p)

    reales = df_crudo[df_crudo.apply(g._fila_real, axis=1)]
    filas = []
    for idx, row in reales.iterrows():
        fila = int(idx) + 2
        geo = g.clasificar_geografia_capital(row["DEPARTAMENTO"], row["LOCALIDAD / SUBCIRCUITO"], row["AUDIENCIA  (Barrios o publicos) "], ref)
        entradas = por_fila.get(fila, [])
        entra = bool(entradas)
        destinos = sorted({p["destino_tablero"] for p in entradas})
        motivo = ""
        if not entra:
            motivo = (geo["motivo"] or "Sin entradas") if geo["destino"] in ("excluido", "pendiente") else "Todas sus entradas colapsaron por duplicado"
        filas.append({
            "fila": fila, "cod": g.limpio(row["COD. CONTENIDO"]), "campania": g.limpio(row["CAMPAÑA"]), "titulo": g.limpio(row["TITULO DE LA PIEZA"]),
            "depto": g.limpio(row["DEPARTAMENTO"]), "loc": g.limpio(row["LOCALIDAD / SUBCIRCUITO"]), "estado_crudo": g.limpio(row["ESTADO"]),
            "inicio": _a_fecha(row["INICIO "]), "fin": _a_fecha(row["FIN"]),
            "tipo_geo": tipo_geografia(row["LOCALIDAD / SUBCIRCUITO"]), "destino": " + ".join(destinos) if destinos else geo["destino"],
            "entra": "SI" if entra else "NO", "motivo": motivo, "entradas": len(entradas),
            "estado_tablero": entradas[0]["estado"] if entradas else g.estado_pauta_desde_crudo(row["ESTADO"]),
            "dirigida": 1 if any(p["destino_tablero"] in ("barrio", "circuito_sin_barrio") for p in entradas) else 0,
            "p": entradas,
        })

    # Piezas únicas activas por fecha: primera fila de cada clave (alcance Capital y específico por separado)
    marca_unica = {f: {} for f in args.fechas}
    for f in args.fechas:
        vistos = set()
        for r in filas:
            r.setdefault("unica", {})
            ok = 0
            for p in r["p"]:
                if not activa(p, f):
                    continue
                alcance = "capital" if p["destino_tablero"] == "toda_capital" else "especifica"
                k = (alcance, clave_pieza(p))
                if k not in vistos:
                    vistos.add(k)
                    ok = 1
            r["unica"][f] = ok

    wb = Workbook()
    ws = wb.active
    ws.title = "Verificación"
    cab = ["Fila del crudo", "COD. CONTENIDO", "CAMPAÑA", "TITULO DE LA PIEZA", "DEPARTAMENTO", "LOCALIDAD / SUBCIRCUITO", "ESTADO (crudo)", "INICIO", "FIN",
           "Tipo de geografía", "Destino en el tablero", "¿Entra?", "Motivo si no entra", "Entradas generadas", "Estado en el tablero", "Dirigida (barrio o circuito)"]
    for f in args.fechas:
        d, m = f[8:10], f[5:7]
        cab += [f"Activa {d}/{m}", f"Activa dirigida {d}/{m}", f"Pieza única activa {d}/{m}"]
    ws.append(cab)
    for r in filas:
        linea = [r["fila"], r["cod"], r["campania"], r["titulo"], r["depto"], r["loc"], r["estado_crudo"], r["inicio"], r["fin"], r["tipo_geo"], r["destino"], r["entra"],
                 r["motivo"], r["entradas"], r["estado_tablero"], r["dirigida"]]
        for f in args.fechas:
            act = 1 if any(activa(p, f) for p in r["p"]) else 0
            dirig = 1 if any(activa(p, f) and p["destino_tablero"] in ("barrio", "circuito_sin_barrio") for p in r["p"]) else 0
            linea += [act, dirig, r["unica"][f]]
        ws.append(linea)
        for col in (8, 9):
            ws.cell(ws.max_row, col).number_format = "dd/mm/yyyy"
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="F2F2F2")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate([9, 13, 26, 44, 30, 34, 12, 11, 11, 30, 26, 8, 40, 10, 12, 10] + [11] * (3 * len(args.fechas)), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 48

    rs = wb.create_sheet("Resumen")
    rs.append(["Qué se cuenta", "Valor", "Unidad"])
    rs.append(["Filas reales del crudo", len(filas), "filas"])
    rs.append(["Filas que entran al tablero", sum(1 for r in filas if r["entra"] == "SI"), "filas"])
    rs.append(["Entradas del tablero (después de expandir y deduplicar)", sum(r["entradas"] for r in filas), "entradas"])
    rs.append(["Entradas colapsadas por duplicado", stats["entradas_colapsadas_por_dedup"], "entradas"])
    for k, v in Counter(r["estado_tablero"] for r in filas if r["entra"] == "SI").most_common():
        rs.append([f"Filas que entran con estado {k}", v, "filas"])
    for f in args.fechas:
        d, m = f[8:10], f[5:7]
        rs.append([f"Filas activas al {d}/{m}", sum(1 for r in filas if any(activa(p, f) for p in r["p"])), "filas"])
        rs.append([f"Filas activas dirigidas al {d}/{m}", sum(1 for r in filas if any(activa(p, f) and p["destino_tablero"] in ("barrio", "circuito_sin_barrio") for p in r["p"])), "filas"])
        rs.append([f"Entradas activas al {d}/{m}", sum(1 for r in filas for p in r["p"] if activa(p, f)), "entradas"])
        rs.append([f"Piezas únicas activas al {d}/{m}", sum(r["unica"][f] for r in filas), "piezas únicas"])
    rs.append([])
    rs.append(["Motivo si no entra", "Filas", ""])
    for k, v in Counter(re.sub(r"(localidad|departamento|audiencia) = .*$", r"\1", re.sub(r"^(PENDIENTE: [^(]*?)\s*\(.*$", r"\1", r["motivo"])) for r in filas if r["entra"] == "NO").most_common():
        rs.append([k, v, "filas"])
    for c in rs[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="F2F2F2")
    rs.column_dimensions["A"].width = 64
    wb.save(args.salida)
    print(f"Escrito {args.salida} ({len(filas)} filas)")


if __name__ == "__main__":
    main()
