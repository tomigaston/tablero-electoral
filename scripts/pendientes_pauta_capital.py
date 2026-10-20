#!/usr/bin/env python3
"""
Lista de lo que el área tiene que confirmar para que entre al tablero (Excel de tres hojas).

Política: solo se aplica lo confirmado; lo que no se entiende o no respeta la regla queda
afuera. Este script junta, para una carga:

  - Alias pendientes: correcciones de nombres de barrio que el generador conoce pero NO aplica
    (alias_barrios.json -> pendientes_de_confirmar), con cuántos ítems afectaría cada una.
  - Ítems sin resolver: nombres de las audiencias que no se pudieron asignar a un barrio.
  - Reglas a confirmar: reglas que se aplican hoy y que el área todavía no confirmó.

Las filas pendientes (alcance que no se puede decidir con el crudo tal como viene) NO se
piden al área: el crudo se toma tal cual; si una fila no cumple las reglas del manual, queda
afuera y está en el archivo de verificación con su motivo.

Uso:
    python3 scripts/pendientes_pauta_capital.py
    (lee Contexto/Pauta RRSS_GLOBAL.xlsx y escribe Contexto/Pendientes_de_confirmacion.xlsx)

Cómo confirmar: un alias confirmado pasa de "pendientes_de_confirmar" a "alias" en
scripts/config_pauta/alias_barrios.json (commit de código con un caso de test).
"""
import argparse
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
import generar_pauta_capital as g  # noqa: E402

REGLAS_A_CONFIRMAR = [
    ("Ítem de lista sin letra y con el nombre del circuito (ej. '5 Altamira')", "Va a circuito sin barrio: apunta al circuito entero. Con letra ('5 G San Vicente') se busca como barrio.", "Geografía → Audiencias con números de circuito"),
    ("Ítem de lista numerado que no es un barrio reconocido ni el nombre del circuito", "Queda afuera; la fila entra con los demás ítems.", "Geografía → Audiencias con números de circuito"),
    ("Audiencia que solo repite el texto de LOCALIDAD (ej. '007B - ALTA CORDOBA, CÓRDOBA CAPITAL, ARGENTINA sin advantage+…')", "El nombre es el del circuito, no un barrio: la fila va a circuito sin barrio.", "Geografía → Barrios en la columna AUDIENCIA"),
    ("'Toda la provincia excluyendo X' con departamento 'Provincia de Córdoba'", "Va a toda la Capital.", "Geografía → Reglas finas"),
    ("Localidad 'CÓRDOBA CAPITAL…' con audiencia genérica ('Hombres y mujeres de 30 a 65 años', 'Cuadrante 4')", "Va a toda la Capital.", "Geografía → tabla, regla 2"),
    ("Código de circuito en LOCALIDAD y una audiencia o departamento de otra ciudad", "El código manda: la fila va al circuito.", "Geografía → Reglas finas"),
    ("Audiencia que nombra un barrio que no se reconoce (ej. 'Barrio Felipe Varela') y ningún otro barrio", "La fila queda pendiente: no entra. No se asume toda la Capital ni el circuito.", "Geografía → Reglas finas"),
    ("Plataforma: todo lo que no se llama exactamente 'Meta' (Youtube, Google Ads, Programática, DV 360, vacío)", "Se muestra como GOOGLE y se leen las columnas de métricas de Google.", "Plataforma y métricas"),
]


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--crudo", default=str(g.CRUDO_DEFECTO), help="por defecto Contexto/Pauta RRSS_GLOBAL.xlsx")
    ap.add_argument("--salida", default=str(g.ROOT / "Contexto" / "Pendientes_de_confirmacion.xlsx"), help="por defecto Contexto/Pendientes_de_confirmacion.xlsx (no se versiona)")
    args = ap.parse_args(argv)

    ref = g.cargar_referencias()
    df = pd.read_excel(args.crudo, sheet_name="GLOBAL_ImportRange")
    _, _, avisos = g.generar(df, ref)

    # Ítems sin resolver, agrupados por texto
    items = Counter()
    sugerencia = {}
    filas_de = defaultdict(list)
    for fila, aviso in avisos:
        m = re.match(r"Sin match: (.*?)(?: \(alias pendiente de confirmar: (.*)\))?$", aviso)
        if not m:
            filas_de[("otro", aviso)].append(fila)
            items[("otro", aviso)] += 1
            continue
        nombre = m.group(1)
        k = g.normalizar_sin_puntos(nombre)
        items[(k, nombre)] += 1
        filas_de[(k, nombre)].append(fila)
        if m.group(2):
            sugerencia[(k, nombre)] = m.group(2)

    wb = Workbook()

    def hoja(titulo, cab, filas, anchos):
        ws = wb.create_sheet(titulo)
        ws.append(cab)
        for f in filas:
            ws.append(list(f))
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="F2F2F2")
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, w in enumerate(anchos, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    wb.remove(wb.active)
    pend_alias = ref["alias_pendientes"]
    afectados = Counter()
    for (k, _), n in items.items():
        if k in pend_alias:
            afectados[k] += n
    hoja("Alias pendientes", ["Texto en la audiencia (normalizado)", "Barrio sugerido", "Ítems afectados en esta carga", "¿Se confirma? (SI / NO / otro barrio)"],
         sorted(((k, v, afectados.get(k, 0), "") for k, v in pend_alias.items()), key=lambda x: -x[2]), [44, 36, 18, 30])
    hoja("Ítems sin resolver", ["Texto en la audiencia", "Veces en esta carga", "Alias pendiente sugerido", "Filas del crudo (primeras 8)", "Barrio correcto (completa el área)"],
         [(n if k != "otro" else n, c, sugerencia.get((k, n), ""), ", ".join(map(str, sorted(set(filas_de[(k, n)]))[:8])), "") for (k, n), c in items.most_common()], [48, 12, 34, 40, 34])
    hoja("Reglas a confirmar", ["Caso", "Qué hace hoy el tablero", "Dónde está en el manual", "¿Se confirma? (SI / NO)", "Comentario del área"],
         [r + ("", "") for r in REGLAS_A_CONFIRMAR], [70, 70, 40, 22, 40])
    wb.save(args.salida)
    print(f"Escrito {args.salida}: {len(pend_alias)} alias pendientes, {len(items)} ítems distintos sin resolver ({sum(items.values())} apariciones), {len(REGLAS_A_CONFIRMAR)} reglas a confirmar")


if __name__ == "__main__":
    main()
