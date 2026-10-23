#!/usr/bin/env python3
"""
Genera pautas_barrios.json (pauta digital de Córdoba Capital) directo desde el
crudo Pauta RRSS_GLOBAL.xlsx (hoja GLOBAL_ImportRange). Mismo modelo que el
tablero del interior (Tablero comunicacion/scripts/generar_datos.py): un script
versionado, reproceso siempre completo desde el crudo, sin maestro intermedio.

Uso (igual que en el tablero del interior: se pisa el crudo y se corre):
    1) Copiar el crudo nuevo a  Contexto/Pauta RRSS_GLOBAL.xlsx  (pisa el anterior).
    2) python3 scripts/generar_pauta_capital.py
       Lee Contexto/Pauta RRSS_GLOBAL.xlsx y escribe public/data/pautas_barrios.json y
       public/data/Filas_no_procesadas_pauta_capital.xlsx (las filas que no entran, con su
       motivo, para el área). Si el resumen del documento no cuadra con el JSON, se detiene.

Opciones:
    --solo-verificar        imprime la verificación sin escribir ningún archivo
    --crudo <ruta>          otro crudo (por defecto Contexto/Pauta RRSS_GLOBAL.xlsx)
    --salida <ruta>         otro destino (por defecto public/data/pautas_barrios.json)
    --no-procesadas <ruta>  otro destino del documento (por defecto, junto al JSON)
    --previo <ruta>         JSON anterior para el freno de bajas (por defecto, el mismo que se pisa)
    --aceptar-bajas         continuar aunque desaparezcan piezas (solo con confirmación)

Reglas (ver docs/pauta-capital-reglas.md cuando exista, y el manual metodológico):
  - Fila real: CAMPAÑA o TITULO DE LA PIEZA no vacíos (igual que el interior).
  - ESTADO directo del crudo: ACTIVO->Vigente, INACTIVO->Finalizada,
    OBSERVADO->Observada, PROGRAMADO->Programada, vacío->S/D, otro->tal cual.
    Nunca se descarta una fila por su estado.
  - Fechas: solo INICIO y FIN; sin respaldo en INICIO/FIN INFORMADO.
  - Geografía: ver clasificar_geografia_capital().
  - Duplicados: ver clave_duplicado() (CAMPAÑA, TITULO, PLATAFORMA, destino, INICIO, FIN,
    LINK MATERIAL, TIPO DE MATERIAL y AUDIENCIA, normalizados); gana la FECHA DE ENVÍO más
    reciente, luego ESTADO=ACTIVO, luego la primera. COD. CONTENIDO es solo informativo:
    no decide si una fila es válida ni si es duplicada (solo una agencia lo carga).
  - Una fila que nombra N barrios genera N entradas (una por barrio), con las
    métricas completas (no se prorratean). Lo mismo una seccional: una entrada por
    subcircuito (ver entradas_circuito_sin_barrio()).
"""
import argparse
import json
import re
import sys
import unicodedata
from collections import Counter
from datetime import date, datetime
from pathlib import Path

import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "public" / "data"
CRUDO_DEFECTO = ROOT / "Contexto" / "Pauta RRSS_GLOBAL.xlsx"
SALIDA_DEFECTO = DATA / "pautas_barrios.json"
NO_PROCESADAS = "Filas_no_procesadas_pauta_capital.xlsx"  # se escribe junto al JSON
CONFIG = Path(__file__).resolve().parent / "config_pauta"

CODIGO_CIRCUITO_RE = re.compile(r"^\d{3,4}[A-Za-z]?\s*-")  # igual que el interior
ITEM_CIRCUITO_RE = re.compile(r"^\s*(\d{1,2})\s*([A-Za-z])?\s+(\S.*?)\s*$")
TODAS = "TODAS LAS LOCALIDADES"
# Los 26 departamentos de la provincia (normalizados); se usan para reconocer cuando una audiencia nombra otro.
DEPARTAMENTOS_CORDOBA = ['CALAMUCHITA', 'CAPITAL', 'COLON', 'CRUZ DEL EJE', 'GENERAL ROCA', 'GENERAL SAN MARTIN', 'ISCHILIN', 'JUAREZ CELMAN', 'MARCOS JUAREZ', 'MINAS', 'POCHO', 'PRESIDENTE ROQUE SAENZ PENA', 'PUNILLA', 'RIO CUARTO', 'RIO PRIMERO', 'RIO SECO', 'RIO SEGUNDO', 'SAN ALBERTO', 'SAN JAVIER', 'SAN JUSTO', 'SANTA MARIA', 'SOBREMONTE', 'TERCERO ARRIBA', 'TOTORAL', 'TULUMBA', 'UNION']
ESTADOS_MAPEADOS = {"Vigente", "Finalizada", "Observada", "Programada", "S/D"}


# ---------------------------------------------------------------------------
# Helpers (mismo comportamiento que generar_datos.py del interior)
# ---------------------------------------------------------------------------

def normalizar(s):
    """Mayúsculas, sin espacios de más, sin tildes."""
    if s is None or (isinstance(s, float) and pd.isna(s)):
        return ""
    s = str(s).strip().upper()
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", s)


def normalizar_sin_puntos(s):
    """Como normalizar() pero sin puntos ni signos: 'A.T. E.' y 'A.T. E' coinciden."""
    s = normalizar(s).replace(".", " ")
    s = re.sub(r"[^A-Z0-9 ]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def _a_timestamp(v):
    """Timestamp/datetime/date/str -> pd.Timestamp o None. Igual que el interior,
    salvo que un número suelto (serial de Excel sin formato de fecha) NO se
    interpreta: pd.to_datetime(46295) daría 1970-01-01; acá es None."""
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(v, pd.Timestamp):
        return v
    if isinstance(v, (datetime, date)):
        return pd.Timestamp(v)
    if isinstance(v, (int, float)):
        return None
    ts = pd.to_datetime(v, errors="coerce", dayfirst=False)
    return None if pd.isna(ts) else ts


def fecha_iso(v):
    ts = _a_timestamp(v)
    return None if ts is None else ts.strftime("%Y-%m-%d")


def limpio(v):
    """NaN -> None; '-' y vacío -> None; fechas -> ISO; numpy -> nativo."""
    if v is None:
        return None
    if isinstance(v, float) and pd.isna(v):
        return None
    if isinstance(v, (pd.Timestamp, datetime, date)):
        return fecha_iso(v)
    if isinstance(v, str):
        s = v.strip()
        return None if s in ("", "-") else s
    if isinstance(v, (int, bool)):
        return v
    if hasattr(v, "item"):
        return v.item()
    if isinstance(v, float):
        return v
    return str(v)


def limpiar_link(v):
    """None si la celda es '-'/vacía o no contiene 'http'."""
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    s = str(v).strip()
    if s in ("", "-") or "http" not in s.lower():
        return None
    return s


def numero(v):
    """Métricas: '-', vacío y errores de Excel (#DIV/0!, #ERROR!...) -> None."""
    v = limpio(v)
    if isinstance(v, str) and v.startswith("#"):
        return None
    return v


def estado_pauta_desde_crudo(valor_estado_crudo):
    """Estado expuesto: directo de la columna ESTADO, sin inventar categorías ni
    descartar filas. ACTIVO->Vigente, INACTIVO->Finalizada, OBSERVADO->Observada,
    PROGRAMADO->Programada, vacío->S/D, cualquier otro valor tal cual."""
    valor = limpio(valor_estado_crudo)
    if valor is None:
        return "S/D"
    return {
        "ACTIVO": "Vigente",
        "INACTIVO": "Finalizada",
        "OBSERVADO": "Observada",
        "PROGRAMADO": "Programada",
    }.get(normalizar(valor), valor)


def _fila_real(row):
    """Fila real: CAMPAÑA o TITULO DE LA PIEZA no vacíos."""
    return limpio(row["CAMPAÑA"]) is not None or limpio(row["TITULO DE LA PIEZA"]) is not None


def normalizar_circuito(codigo):
    """'010A' -> '10A', '0010' -> '10', '001' -> '1' (CircuitoKey sin ceros a la izquierda)."""
    return re.sub(r"^0+(?=\d)", "", str(codigo or "").strip().upper())


# ---------------------------------------------------------------------------
# Referencias y configuración
# ---------------------------------------------------------------------------

def cargar_referencias(data_dir=DATA, config_dir=CONFIG):
    geo = json.load(open(data_dir / "barrios_geo.json", encoding="utf8"))
    oficiales = sorted({f["properties"]["nombre_barrio"] for f in geo["features"] if f["properties"].get("nombre_barrio")})
    por_norm = {normalizar_sin_puntos(b): b for b in oficiales}

    circuitos = json.load(open(data_dir / "circuitos_completo.json", encoding="utf8"))
    circuito_nombre = {str(c["CircuitoKey"]): normalizar(c["cirn_nom"]) for c in circuitos}
    circuito_nombre_np = {k: normalizar_sin_puntos(v) for k, v in circuito_nombre.items()}
    nombre_a_circuito = {v: k for k, v in circuito_nombre.items()}
    # Seccional (número sin letra) -> sus subcircuitos con letra: '14' -> ['14A', …, '14Q'].
    subcircuitos = {}
    for k in circuito_nombre:
        m = re.match(r"^(\d+)[A-Z]$", k)
        if m:
            subcircuitos.setdefault(m.group(1), []).append(k)
    subcircuitos = {s: sorted(v) for s, v in subcircuitos.items()}

    alias_cfg = json.load(open(config_dir / "alias_barrios.json", encoding="utf8"))
    alias = {normalizar_sin_puntos(k): v for k, v in alias_cfg["alias"].items()}
    for destino in alias.values():
        if normalizar_sin_puntos(destino) not in por_norm:
            raise SystemExit(f"alias_barrios.json: '{destino}' no es un barrio oficial de barrios_geo.json")
    pendientes = {normalizar_sin_puntos(k): v for k, v in alias_cfg.get("pendientes_de_confirmar", {}).items()}
    for destino in pendientes.values():
        if normalizar_sin_puntos(destino) not in por_norm:
            raise SystemExit(f"alias_barrios.json (pendientes): '{destino}' no es un barrio oficial de barrios_geo.json")
    descartar = {normalizar_sin_puntos(x) for x in alias_cfg.get("descartar", [])}

    reglas = json.load(open(config_dir / "decisiones_filas.json", encoding="utf8"))["reglas"]
    return {
        "oficiales": oficiales, "por_norm": por_norm, "circuito_nombre": circuito_nombre,
        "nombre_a_circuito": nombre_a_circuito, "circuito_nombre_np": circuito_nombre_np, "subcircuitos": subcircuitos, "alias": alias, "alias_pendientes": pendientes, "descartar": descartar, "reglas": reglas,
    }


# ---------------------------------------------------------------------------
# Barrios en AUDIENCIA
# ---------------------------------------------------------------------------

def _nombres_buscables(ref):
    """Diccionario texto-normalizado -> barrio oficial (oficiales + alias), más largo primero."""
    nombres = dict(ref["por_norm"])
    for k, v in ref["alias"].items():
        nombres[k] = v
    return sorted(nombres.items(), key=lambda kv: -len(kv[0])), nombres


def extraer_barrios(texto, ref):
    """Greedy longest-match de nombres de barrio (oficiales y alias) dentro de un
    texto libre. Devuelve la lista de barrios oficiales en orden de aparición, sin
    repetir. Los nombres de la lista 'descartar' se consumen pero no se devuelven."""
    t = " " + normalizar_sin_puntos(texto) + " "
    ordenados, _ = _nombres_buscables(ref)
    hallados = []  # (posicion, barrio)
    for nombre, oficial in ordenados:
        if not nombre:
            continue
        patron = re.compile(r"(?<![A-Z0-9])" + re.escape(nombre) + r"(?![A-Z0-9])")
        for m in patron.finditer(t):
            hallados.append((m.start(), oficial))
            t = t[:m.start()] + " " * (m.end() - m.start()) + t[m.end():]
    for d in ref["descartar"]:
        patron = re.compile(r"(?<![A-Z0-9])" + re.escape(d) + r"(?![A-Z0-9])")
        t = patron.sub(lambda m: " " * len(m.group(0)), t)
    vistos, resultado = set(), []
    for _, b in sorted(hallados):
        if b not in vistos:
            vistos.add(b)
            resultado.append(b)
    return resultado


def resolver_nombre_barrio(nombre, ref):
    """Un ítem de una lista -> barrio oficial, o None. Orden: barrio oficial exacto, alias
    CONFIRMADO. No hay coincidencia aproximada: lo que no se entiende queda afuera."""
    n = normalizar_sin_puntos(nombre)
    if n in ref["descartar"]:
        return "__DESCARTAR__"
    if n in ref["por_norm"]:
        return ref["por_norm"][n]
    if n in ref["alias"]:
        return ref["alias"][n]
    return None


def aviso_sin_match(nombre, ref):
    """Texto del aviso de un ítem que quedó afuera; sugiere el alias pendiente si existe."""
    n = normalizar_sin_puntos(nombre)
    sugerido = ref["alias_pendientes"].get(n)
    extra = f" (alias pendiente de confirmar: {sugerido})" if sugerido else ""
    return f"Sin match: {nombre}{extra}"


def segmentos_sin_match(aud, ref):
    """Audiencia en forma de lista (3 o más segmentos separados por coma o salto de línea):
    segmentos cortos sin ningún barrio reconocido. Se informan; no se asignan."""
    partes = [p.strip() for p in re.split(r"[\r\n]+|,", str(aud)) if p.strip()]
    if len(partes) < 3:
        return []
    ruido = {"CORDOBA", "CORDOBA CAPITAL", "CAPITAL", "ARGENTINA", "PROVINCIA DE CORDOBA", "DEPARTAMENTO CAPITAL"}
    return [p for p in partes if len(p.split()) <= 6 and normalizar(p) not in ruido
            and not re.search(r"\d|A[ÑN]OS|ADVANTAGE|\bKM\b", normalizar(p)) and not extraer_barrios(p, ref)]


def menciona_barrio(aud):
    """La audiencia habla de 'barrio/barrios' explícitamente."""
    return bool(re.search(r"\bBARRIOS?\b", normalizar(aud)))


def menciona_otro_departamento(aud):
    """Nombra un departamento de Córdoba que no es Capital (por ejemplo 'Departamento Punilla')."""
    t = normalizar(aud).replace("DEPATAMENTO", "DEPARTAMENTO")
    for d in DEPARTAMENTOS_CORDOBA:
        if d != "CAPITAL" and re.search(r"\bDEPARTAMENTO\s+(?:DE\s+)?" + re.escape(d) + r"\b", t):
            return True
    return False


def _dividir_items(aud):
    partes = re.split(r"[\r\n]+|,|\s+y\s+(?=\d)", str(aud))
    return [p.strip() for p in partes if p and p.strip()]


def items_con_circuito(aud):
    """Si la audiencia es una lista de ítems con número de circuito ('14 B Argüello
    Primero, 9 B Las Margaritas, 5 Altamira'), devuelve [(CircuitoKey, nombre, tiene_letra)]; si no, None."""
    partes = _dividir_items(aud)
    if len(partes) < 2:
        return None
    items = []
    for p in partes:
        m = ITEM_CIRCUITO_RE.match(p)
        if m:
            items.append((m.group(1) + (m.group(2) or "").upper(), m.group(3), bool(m.group(2))))
    if len(items) * 3 < len(partes) * 2:  # al menos 2/3 de los ítems con número
        return None
    return items


# ---------------------------------------------------------------------------
# Geografía de Capital
# ---------------------------------------------------------------------------

# Una seccional que llega como tal (código sin letra) también cubre su propio circuito sin
# letra ('14' ARGUELLO, que es uno de los 119), además de 14A…14Q.
SECCIONAL_INCLUYE_CIRCUITO_PROPIO = True


def entradas_circuito_sin_barrio(circuito, ref):
    """Circuito con letra -> una entrada. Seccional (sin letra) -> una entrada por cada
    subcircuito de circuitos_completo.json (y la del propio circuito según
    SECCIONAL_INCLUYE_CIRCUITO_PROPIO); una seccional sin subcircuitos (1, 2, 3) queda en sí misma."""
    subs = ref["subcircuitos"].get(circuito, []) if circuito.isdigit() else []
    if not subs:
        destinos = [circuito]
    else:
        destinos = ([circuito] if SECCIONAL_INCLUYE_CIRCUITO_PROPIO else []) + subs
    return [{"destino": "circuito_sin_barrio", "circuito": c, "barrio": None} for c in destinos]

def _depto_candidato(depto_raw):
    """'DEPARTAMENTO DE SAN JUSTO, PROVINCIA…' -> 'SAN JUSTO'; sin prefijo DEPARTAMENTO -> None."""
    if depto_raw is None or (isinstance(depto_raw, float) and pd.isna(depto_raw)):
        return None
    primero = str(depto_raw).split(",")[0].strip()
    m = re.match(r"^DEPARTAMENTO\s+(DE\s+)?(.+)$", primero, flags=re.IGNORECASE)
    return m.group(2).strip() if m else None


def _aplica_regla_fila(aud, ref):
    texto = str(aud or "")
    for r in ref["reglas"]:
        c = r.get("coincide", {})
        if "audiencia_empieza_con" in c and texto.strip().startswith(c["audiencia_empieza_con"]):
            return r
    return None


def clasificar_geografia_capital(depto_raw, loc_raw, aud_raw, ref):
    """Devuelve un dict con:
        destino: 'excluido' | 'pendiente' | 'toda_capital' | 'entradas'
        motivo: texto para observaciones (o None)
        entradas: [{'destino': 'barrio'|'circuito_sin_barrio', 'circuito': str|None, 'barrio': str|None}]
        avisos: [texto]  (barrios que no se pudieron resolver; la fila entra igual con el resto)
    Orden de evaluación:
      1) DEPARTAMENTO sin 'CÓRDOBA' -> excluido (Santa Fe, Entre Ríos…).
      2) Código de circuito en LOCALIDAD / SUBCIRCUITO -> manda el código; con barrios en
         AUDIENCIA, una entrada por barrio; sin barrios, circuito_sin_barrio.
      3) 'TODAS LAS LOCALIDADES': decide DEPARTAMENTO — Capital -> toda_capital; otro
         departamento u otra provincia -> excluido; 'PROVINCIA DE CÓRDOBA' sin departamento
         -> decide AUDIENCIA (menciona Córdoba/Capital/Provincia -> toda_capital; nombra
         otro departamento -> excluido; sin referencia geográfica -> pendiente).
      4) 'CÓRDOBA CAPITAL, DEPARTAMENTO CAPITAL…': decide AUDIENCIA (lista con números de
         circuito, barrios nombrados, o toda_capital). Si la audiencia nombra otro
         departamento o un barrio que no se pudo resolver -> pendiente (queda afuera).
    Política: solo se aplica lo confirmado por el área. Lo que no se entiende o no respeta la
    regla queda afuera (pendiente o ítem omitido) y se informa en la corrida; no hay
    coincidencia aproximada ni alias sin confirmar.
      5) Cualquier otra localidad -> excluido.
    """
    res = {"destino": None, "motivo": None, "entradas": [], "avisos": []}
    n_depto = normalizar(depto_raw)
    loc = None if (loc_raw is None or (isinstance(loc_raw, float) and pd.isna(loc_raw))) else str(loc_raw).strip()
    aud = limpio(aud_raw) or ""

    regla = _aplica_regla_fila(aud, ref)
    if regla:
        res.update(destino=regla["destino"], motivo=regla["motivo"])
        return res

    if "CORDOBA" not in n_depto:
        res.update(destino="excluido", motivo=f"Fuera de Capital: departamento = {limpio(depto_raw) or 'vacío'}")
        return res

    if loc and CODIGO_CIRCUITO_RE.match(loc):
        circuito = normalizar_circuito(loc.split("-")[0])
        # Muchas filas traen en AUDIENCIA el texto de LOCALIDAD repetido ("007B - ALTA CORDOBA,
        # CÓRDOBA CAPITAL, ARGENTINA sin advantage+…"): ese nombre es el del circuito, no un barrio.
        aud_sin_loc = re.sub(re.escape(loc), " ", aud, flags=re.IGNORECASE)
        barrios = extraer_barrios(aud_sin_loc, ref)
        if barrios:
            res["entradas"] = [{"destino": "barrio", "circuito": circuito, "barrio": b} for b in barrios]
            res["avisos"] += [aviso_sin_match(x, ref) for x in segmentos_sin_match(aud_sin_loc, ref)]
        elif menciona_barrio(aud_sin_loc):
            res.update(destino="pendiente", motivo=f"PENDIENTE: la audiencia nombra un barrio que no se pudo resolver ({aud_sin_loc.strip()})")
            return res
        else:
            res["entradas"] = entradas_circuito_sin_barrio(circuito, ref)
        res["destino"] = "entradas"
        return res

    if loc == TODAS:
        depto = _depto_candidato(depto_raw)
        if depto:
            if normalizar(depto) == "CAPITAL":
                res.update(destino="toda_capital")
            else:
                res.update(destino="excluido", motivo=f"Fuera de Capital: localidad = {TODAS}; departamento = {limpio(depto_raw)}")
            return res
        # 'PROVINCIA DE CÓRDOBA, ARGENTINA' sin departamento puntual: decide AUDIENCIA
        n_aud = normalizar(aud)
        if not n_aud:
            res.update(destino="pendiente", motivo="PENDIENTE: alcance geográfico ambiguo (sin audiencia)")
        elif re.search(r"\bDEPARTAMENTO\b", n_aud):
            res.update(destino="excluido", motivo=f"Fuera de Capital: localidad = {TODAS}; audiencia = {aud}")
        elif re.search(r"\b(CORDOBA|CAPITAL|PROVINCIA)\b", n_aud):
            res.update(destino="toda_capital")
        else:
            res.update(destino="pendiente", motivo=f"PENDIENTE: alcance geográfico ambiguo ({aud})")
        return res

    if loc:
        candidata = normalizar(loc.split(",")[0])
        if candidata == "CORDOBA CAPITAL":
            items = items_con_circuito(aud)
            if items:
                entradas = []
                for circuito, nombre, tiene_letra in items:
                    if circuito not in ref["circuito_nombre"]:
                        res["avisos"].append(f"circuito '{circuito}' no existe: {nombre}")
                        continue
                    # Ítem SIN letra cuyo nombre es el del propio circuito ('5 Altamira'): apunta al
                    # seccional entera, no a un barrio -> circuito_sin_barrio en cada subcircuito.
                    # Con letra ('5 G San Vicente') el nombre se busca como barrio.
                    if not tiene_letra and normalizar_sin_puntos(nombre) == ref["circuito_nombre_np"][circuito]:
                        entradas += entradas_circuito_sin_barrio(circuito, ref)
                        continue
                    barrio = resolver_nombre_barrio(nombre, ref)
                    if barrio == "__DESCARTAR__":
                        continue
                    if barrio:
                        entradas.append({"destino": "barrio", "circuito": circuito, "barrio": barrio})
                    else:
                        res["avisos"].append(aviso_sin_match(nombre, ref))
                res["entradas"] = entradas
                res["destino"] = "entradas" if entradas else "pendiente"
                if not entradas:
                    res["motivo"] = "PENDIENTE: ningún ítem de la audiencia se pudo asignar"
                return res
            barrios = extraer_barrios(aud, ref)
            if barrios:
                res["entradas"] = [{"destino": "barrio", "circuito": None, "barrio": b} for b in barrios]
                res["avisos"] += [aviso_sin_match(x, ref) for x in segmentos_sin_match(aud, ref)]
                res["destino"] = "entradas"
            elif menciona_otro_departamento(aud):
                res.update(destino="pendiente", motivo=f"PENDIENTE: la audiencia nombra otro departamento ({aud})")
            elif menciona_barrio(aud):
                res.update(destino="pendiente", motivo=f"PENDIENTE: la audiencia nombra un barrio que no se pudo resolver ({aud})")
            else:
                res["destino"] = "toda_capital"
            return res
        res.update(destino="excluido", motivo=f"Fuera de Capital: localidad = {loc}")
        return res

    res.update(destino="pendiente", motivo="PENDIENTE: LOCALIDAD / SUBCIRCUITO vacío")
    return res


# ---------------------------------------------------------------------------
# Duplicados
# ---------------------------------------------------------------------------

def _gana_pieza(candidata, actual):
    """True si `candidata` reemplaza a `actual`: 1) FECHA DE ENVÍO más reciente (None es
    la más vieja); 2) en empate, ESTADO=ACTIVO; 3) si sigue el empate, se queda `actual`."""
    ts_c, ts_a = candidata["_fecha_envio_ts"], actual["_fecha_envio_ts"]
    if ts_c != ts_a:
        if ts_a is None:
            return True
        if ts_c is None:
            return False
        return ts_c > ts_a
    if candidata["_estado_activo"] != actual["_estado_activo"]:
        return candidata["_estado_activo"]
    return False


def clave_duplicado(e):
    """CAMPAÑA + TITULO DE LA PIEZA + PLATAFORMA (cruda) + destino + INICIO + FIN + LINK MATERIAL +
    TIPO DE MATERIAL + AUDIENCIA, normalizados (espacios, mayúsculas, tildes). COD. CONTENIDO no
    participa: solo una agencia lo carga."""
    return (normalizar(e["campania"]), normalizar(e["titulo_pieza"]), normalizar(e["plataforma_origen"]),
            normalizar(e["_lugar"]), e["inicio"] or "", e["fin"] or "",
            normalizar(e["material"]["link"]), normalizar(e["material"]["tipo"]), normalizar(e["_audiencia"]))


def deduplicar(entradas):
    """Unifica las entradas con la misma clave_duplicado(); se conserva una según _gana_pieza().
    Las métricas nunca se suman. Devuelve (lista, descartadas, unificadas), con
    unificadas = [(entrada descartada, entrada que quedó)]."""
    grupos, orden = {}, []
    for e in entradas:
        clave = clave_duplicado(e)
        if clave not in grupos:
            grupos[clave] = [e]
            orden.append(clave)
        else:
            grupos[clave].append(e)
    resultado, unificadas = [], []
    for clave in orden:
        ganadora = grupos[clave][0]
        for e in grupos[clave][1:]:
            if _gana_pieza(e, ganadora):
                ganadora = e
        resultado.append(ganadora)
        unificadas += [(e, ganadora) for e in grupos[clave] if e is not ganadora]
    resultado.sort(key=lambda e: e["_orden"])
    return resultado, len(unificadas), unificadas


# ---------------------------------------------------------------------------
# Generación
# ---------------------------------------------------------------------------

def _metricas(row, es_meta):
    inversion = numero(row["Inversión ejecutada "])
    if es_meta:
        return {
            "alcance_o_visualizaciones": numero(row["Impresiones Ejecutada META"]),
            "frecuencia": numero(row["Frecuencia (Meta) Ejecutada"]),
            "costo_unitario": numero(row["CPM Real META"]),
            "inversion": inversion,
        }
    return {
        "alcance_o_visualizaciones": numero(row["Visualizaciones Ejecutadas GOOGLE"]),
        "frecuencia": None,
        "costo_unitario": numero(row["CPV Real GOOGLE"]),
        "inversion": inversion,
    }


def generar(df_crudo, ref, dedup=True):
    """Devuelve (json_salida, estadisticas, avisos_por_fila)."""
    mascara = df_crudo.apply(_fila_real, axis=1)
    df = df_crudo[mascara]
    stats = Counter()
    stats["filas_totales"] = len(df_crudo)
    stats["filas_descartadas"] = len(df_crudo) - len(df)
    stats["filas_reales"] = len(df)
    stats["filas_con_cod_contenido"] = int(df["COD. CONTENIDO"].apply(lambda v: limpio(v) is not None).sum())

    entradas, orden = [], 0
    avisos, detalle = [], []
    for idx, row in df.iterrows():
        geo = clasificar_geografia_capital(row["DEPARTAMENTO"], row["LOCALIDAD / SUBCIRCUITO"], row["AUDIENCIA  (Barrios o publicos) "], ref)
        fila = int(idx) + 2  # número de fila en la hoja del crudo (encabezado en la fila 1)
        stats["filas_por_destino_" + geo["destino"]] += 1
        detalle.append({"fila": fila, "destino": geo["destino"], "motivo": geo["motivo"]})
        for aviso in geo["avisos"]:
            avisos.append((fila, aviso))
        if geo["destino"] in ("excluido", "pendiente"):
            stats["filas_no_entran"] += 1
            continue
        plataforma_origen = limpio(row["PLATAFORMA"])
        es_meta = plataforma_origen == "Meta"
        estado = estado_pauta_desde_crudo(row["ESTADO"])
        inicio_ts, fin_ts = _a_timestamp(row["INICIO "]), _a_timestamp(row["FIN"])
        observaciones = []
        if geo["avisos"]:
            observaciones.append("; ".join(geo["avisos"]))
        base = {
            "fila_origen": fila,
            "cod_contenido": limpio(row["COD. CONTENIDO"]),
            "campania": limpio(row["CAMPAÑA"]),
            "titulo_pieza": limpio(row["TITULO DE LA PIEZA"]),
            "eje": limpio(row["EJE"]),
            "ministerio": limpio(row["MINISTERIO"]),
            "plataforma": "META" if es_meta else "GOOGLE",
            "plataforma_origen": plataforma_origen,
            "objetivo_campana": limpio(row["OBJETIVO DE CAMPAÑA"]),
            "tipo_pauta": limpio(row["TIPO DE PAUTA"]),
            "solicitado_por": limpio(row["SOLICITADO POR"]),
            "material": {
                "link": limpio(row["LINK MATERIAL"]),
                "utms": limpio(row["UTM´S"]),
                "tipo": limpio(row["TIPO DE MATERIAL"]),
                "redes_publicadas": {"ig": limpiar_link(row["IG"]), "fb": limpiar_link(row["FB"]), "yt": limpiar_link(row["YT"])},
            },
            "inicio": fecha_iso(inicio_ts),
            "fin": fecha_iso(fin_ts),
            "duracion_dias": limpio(row["DURACIÓN"]),
            "fecha_envio": fecha_iso(row["FECHA DE ENVÍO"]),
            "estado": estado,
            "metricas_ejecutadas": _metricas(row, es_meta),
        }
        if geo["destino"] == "toda_capital":
            lista = [{"destino": "toda_capital", "circuito": None, "barrio": None}]
        else:
            lista = geo["entradas"]
        for item in lista:
            orden += 1
            circuito = item["circuito"]
            e = dict(base)
            e.update({
                "circuito": circuito,
                "circuito_nombre": ref["circuito_nombre"].get(circuito) if circuito else None,
                "destino_tablero": item["destino"],
                "observaciones": " | ".join(observaciones) if observaciones else None,
                "barrio": item["barrio"],
                "_lugar": item["barrio"] or ("circuito:" + str(circuito) if item["destino"] == "circuito_sin_barrio" else "CAPITAL"),
                "_audiencia": limpio(row["AUDIENCIA  (Barrios o publicos) "]),
                "_fecha_envio_ts": _a_timestamp(row["FECHA DE ENVÍO"]),
                "_estado_activo": normalizar(limpio(row["ESTADO"])) == "ACTIVO",
                "_orden": orden,
            })
            entradas.append(e)
            stats["entradas_antes_de_dedup"] += 1

    entradas, colapsadas, unificadas = deduplicar(entradas) if dedup else (entradas, 0, [])
    stats["entradas_colapsadas_por_dedup"] = colapsadas
    # Para el documento de filas no procesadas: el destino y motivo de cada fila real, y con
    # qué fila se unificó cada entrada descartada (con las métricas de las dos).
    stats["detalle_filas"] = detalle
    stats["filas_que_entran"] = sorted({e["fila_origen"] for e in entradas})
    stats["unificadas"] = [(perdida["fila_origen"], perdida["_lugar"], ganadora["fila_origen"]) for perdida, ganadora in unificadas]
    stats["unificadas_metricas"] = {(perdida["fila_origen"], ganadora["fila_origen"]): (perdida["metricas_ejecutadas"], ganadora["metricas_ejecutadas"])
                                    for perdida, ganadora in unificadas}
    stats["codigos_unicos_con_cod_contenido"] = int(
        df[df["COD. CONTENIDO"].apply(lambda v: limpio(v) is not None)]["COD. CONTENIDO"].nunique()
    )

    salida = {"barrios": {}, "circuitos_sin_barrio": {}, "toda_capital": {"total_pautas": 0, "pautas": []}}
    for e in entradas:
        barrio, destino = e.pop("barrio"), e["destino_tablero"]
        for k in [k for k in e if k.startswith("_")]:
            e.pop(k)
        if destino == "toda_capital":
            salida["toda_capital"]["pautas"].append(e)
        elif destino == "circuito_sin_barrio":
            salida["circuitos_sin_barrio"].setdefault(e["circuito"], {"total_pautas": 0, "pautas": []})["pautas"].append(e)
        else:
            salida["barrios"].setdefault(barrio, {"total_pautas": 0, "pautas": []})["pautas"].append(e)
    for b in ("barrios", "circuitos_sin_barrio"):
        for v in salida[b].values():
            v["total_pautas"] = len(v["pautas"])
    salida["toda_capital"]["total_pautas"] = len(salida["toda_capital"]["pautas"])
    return salida, stats, avisos


def entradas_de(salida):
    for b, v in salida["barrios"].items():
        for p in v["pautas"]:
            yield "barrios", b, p
    for c, v in salida["circuitos_sin_barrio"].items():
        for p in v["pautas"]:
            yield "circuitos_sin_barrio", c, p
    for p in salida["toda_capital"]["pautas"]:
        yield "toda_capital", None, p


# ---------------------------------------------------------------------------
# Documento de filas no procesadas (para el área)
# ---------------------------------------------------------------------------

AUD_COL = "AUDIENCIA  (Barrios o publicos) "
CLAVE_TEXTO = "campaña, título, plataforma, destino, INICIO, FIN, link, tipo de material y audiencia"
BLOQUES_DUPLICADAS = [
    ("Repetición genuina", "Misma pieza cargada dos veces: coincide con otra fila en la clave de duplicado y tiene las mismas métricas."),
    ("Sin métricas", "Coincide con otra fila en la clave de duplicado y ninguna de las dos tiene métricas cargadas."),
    ("A confirmar por el área", "Coincide con otra fila en la clave de duplicado pero las métricas son distintas: puede ser la misma pieza cargada dos veces o dos piezas distintas."),
]


def motivo_pendiente_claro(motivo):
    """Motivo técnico de una fila pendiente -> texto para quien carga la planilla."""
    m = re.match(r"^PENDIENTE: (.*?)(?: \((.*)\))?$", motivo or "", flags=re.S)
    tipo, texto = (m.group(1), m.group(2)) if m else (motivo, None)
    if tipo == "alcance geográfico ambiguo":
        aud = "la AUDIENCIA está vacía" if texto in (None, "sin audiencia") else f"la AUDIENCIA («{texto}») no dice dónde llega"
        return (f"La localidad dice TODAS LAS LOCALIDADES, el departamento es solo «Provincia de Córdoba» y {aud}. "
                "Para que entre, indicar en la AUDIENCIA si llega a toda la Capital, a un barrio o a un circuito, o cargar el departamento.")
    if tipo == "la audiencia nombra otro departamento":
        return f"La localidad es Córdoba Capital pero la AUDIENCIA nombra otro departamento («{texto}»). Corregir la localidad o la audiencia."
    if tipo == "la audiencia nombra un barrio que no se pudo resolver":
        return (f"La AUDIENCIA nombra un barrio que no figura con su nombre oficial ni tiene un alias confirmado («{texto}»). "
                "Escribir el barrio con su nombre oficial.")
    if tipo == "ningún ítem de la audiencia se pudo asignar":
        return ("La AUDIENCIA es una lista con números de circuito, pero ningún ítem es un barrio reconocido ni el nombre de su circuito. "
                "Revisar los nombres de la lista.")
    if tipo == "LOCALIDAD / SUBCIRCUITO vacío":
        return "LOCALIDAD / SUBCIRCUITO está vacía. Cargar la localidad o el circuito."
    return motivo


def motivo_fuera_de_capital(motivo):
    """Motivo técnico de una fila excluida -> categoría para el conteo (no se lista fila por fila)."""
    if motivo.startswith("Fuera de Capital: departamento ="):
        return "El DEPARTAMENTO no es de Córdoba (otra provincia)"
    if motivo.startswith(f"Fuera de Capital: localidad = {TODAS}; departamento ="):
        return "TODAS LAS LOCALIDADES de otro departamento de Córdoba"
    if motivo.startswith(f"Fuera de Capital: localidad = {TODAS}; audiencia ="):
        return "TODAS LAS LOCALIDADES de la provincia y la AUDIENCIA nombra otro departamento"
    if motivo.startswith("Fuera de Capital: localidad ="):
        return "La localidad es otra (no es Córdoba Capital)"
    return f"Regla puntual de decisiones_filas.json: {motivo}"


def _sin_metricas(m):
    return all(v is None for v in m.values())


def clasificar_duplicadas(stats):
    """Filas que no entran porque todas sus entradas se unificaron con otra fila -> lista de
    {fila, conservada (filas), bloque}. El bloque se decide por las métricas de cada pareja."""
    entran = set(stats["filas_que_entran"])
    ganadoras = {}
    for perdida, _, ganadora in stats["unificadas"]:
        if perdida not in entran:
            ganadoras.setdefault(perdida, set()).add(ganadora)
    resultado = []
    for fila in sorted(ganadoras):
        pares = [stats["unificadas_metricas"][(fila, g)] for g in sorted(ganadoras[fila])]
        if all(_sin_metricas(a) and _sin_metricas(b) for a, b in pares):
            bloque = "Sin métricas"
        elif all(a == b for a, b in pares):
            bloque = "Repetición genuina"
        else:
            bloque = "A confirmar por el área"
        resultado.append({"fila": fila, "conservada": sorted(ganadoras[fila]), "bloque": bloque, "metricas": pares})
    return resultado


def _por_que_queda(row_perdida, row_ganadora):
    ta, tb = _a_timestamp(row_perdida["FECHA DE ENVÍO"]), _a_timestamp(row_ganadora["FECHA DE ENVÍO"])
    if ta != tb:
        return "Tiene la FECHA DE ENVÍO más reciente"
    if (normalizar(limpio(row_perdida["ESTADO"])) == "ACTIVO") != (normalizar(limpio(row_ganadora["ESTADO"])) == "ACTIVO"):
        return "Tiene ESTADO = ACTIVO"
    return "Empate en FECHA DE ENVÍO y ESTADO: se conserva la primera de la planilla"


def _fmt_num(v):
    return "s/d" if v is None else f"{float(v):,.0f}".replace(",", ".")


def escribir_no_procesadas(path, df_crudo, stats):
    """Excel para el área: hoja Resumen (cuadre completo), Pendientes y Duplicadas fila por fila,
    y Fuera de Capital solo como conteo por motivo. Devuelve el resumen [(qué, cantidad)]."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    fila_row = lambda f: df_crudo.loc[f - 2]
    detalle = {d["fila"]: d for d in stats["detalle_filas"]}
    entran = set(stats["filas_que_entran"])
    pendientes = [d for d in stats["detalle_filas"] if d["destino"] == "pendiente"]
    fuera = [d for d in stats["detalle_filas"] if d["destino"] == "excluido"]
    duplicadas = clasificar_duplicadas(stats)
    por_bloque = Counter(d["bloque"] for d in duplicadas)
    entradas_finales = stats["entradas_antes_de_dedup"] - stats["entradas_colapsadas_por_dedup"]
    fuera_por_motivo = Counter(motivo_fuera_de_capital(d["motivo"]) for d in fuera)

    resumen = [
        ("Renglones de la planilla (sin el encabezado)", stats["filas_totales"]),
        ("Descartadas por vacías (sin CAMPAÑA ni TITULO DE LA PIEZA)", stats["filas_descartadas"]),
        ("Filas reales", stats["filas_reales"]),
        ("Entran al tablero", len(entran)),
        ("Fuera de Capital", len(fuera)),
        ("Pendientes", len(pendientes)),
        ("Duplicadas (todas sus entradas se unificaron con otra fila)", len(duplicadas)),
    ] + [(f"   Duplicadas: {b.lower()}", por_bloque.get(b, 0)) for b, _ in BLOQUES_DUPLICADAS] + [
        ("Entradas generadas antes de unificar duplicados", stats["entradas_antes_de_dedup"]),
        ("Entradas unificadas por duplicado", stats["entradas_colapsadas_por_dedup"]),
        ("Entradas del tablero (pautas_barrios.json)", entradas_finales),
    ]
    if stats["filas_reales"] != len(entran) + len(fuera) + len(pendientes) + len(duplicadas):
        raise SystemExit("Filas no procesadas: filas reales != entran + fuera de Capital + pendientes + duplicadas")
    if stats["filas_totales"] != stats["filas_descartadas"] + stats["filas_reales"]:
        raise SystemExit("Filas no procesadas: renglones != descartadas por vacías + filas reales")
    if len(fuera) != sum(fuera_por_motivo.values()):
        raise SystemExit("Filas no procesadas: el conteo de fuera de Capital por motivo no suma el total")

    wb = Workbook()
    gris = PatternFill("solid", fgColor="F2F2F2")

    def encabezado(ws, fila=1):
        for c in ws[fila]:
            c.font = Font(bold=True)
            c.fill = gris
            c.alignment = Alignment(wrap_text=True, vertical="top")

    def anchos(ws, lista):
        for i, w in enumerate(lista, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    rs = wb.active
    rs.title = "Resumen"
    rs.append(["Qué se cuenta", "Cantidad"])
    for k, v in resumen:
        rs.append([k, v])
    rs.append([])
    rs.append(["Cuadre", "Resultado"])
    encabezado(rs, rs.max_row)
    rs.append(["Renglones = descartadas por vacías + filas reales", "OK"])
    rs.append(["Filas reales = entran + fuera de Capital + pendientes + duplicadas", "OK"])
    rs.append(["Entradas del tablero = generadas − unificadas = entradas de pautas_barrios.json", "OK"])
    rs.append([])
    rs.append(["Fuera de Capital por motivo (no es un error: solo se informa la cantidad)", "Filas"])
    encabezado(rs, rs.max_row)
    for k, v in fuera_por_motivo.most_common():
        rs.append([k, v])
    rs.append([])
    rs.append(["Carga procesada: hoja GLOBAL_ImportRange de Pauta RRSS_GLOBAL.xlsx. Los números de fila son los de la planilla (encabezado en la fila 1)."])
    rs.append([f"Clave de duplicado: {CLAVE_TEXTO}, sin distinguir mayúsculas, tildes ni espacios. COD. CONTENIDO no se usa para decidir."])
    encabezado(rs)
    anchos(rs, [90, 12])

    cab = ["Fila en la planilla", "CAMPAÑA", "TITULO DE LA PIEZA", "LOCALIDAD / SUBCIRCUITO", "DEPARTAMENTO", "AUDIENCIA", "ESTADO"]

    def datos(f):
        row = fila_row(f)
        return [f, limpio(row["CAMPAÑA"]), limpio(row["TITULO DE LA PIEZA"]), limpio(row["LOCALIDAD / SUBCIRCUITO"]),
                limpio(row["DEPARTAMENTO"]), limpio(row[AUD_COL]), limpio(row["ESTADO"]) or "(vacío)"]

    wp = wb.create_sheet("Pendientes")
    wp.append(cab + ["Motivo", "Motivo técnico (del generador)", "COD. CONTENIDO (informativo)"])
    for d in pendientes:
        wp.append(datos(d["fila"]) + [motivo_pendiente_claro(d["motivo"]), d["motivo"], limpio(fila_row(d["fila"])["COD. CONTENIDO"])])
    encabezado(wp)
    wp.freeze_panes = "B2"
    wp.auto_filter.ref = wp.dimensions
    anchos(wp, [10, 26, 40, 34, 30, 40, 12, 70, 50, 16])

    wd = wb.create_sheet("Duplicadas")
    wd.append(cab + ["Bloque", "Se unifica con la fila (queda en el tablero)", "Impresiones / visualizaciones de esta fila",
                     "Impresiones / visualizaciones de la fila que queda", "Por qué queda la otra fila", "Motivo", "COD. CONTENIDO (informativo)"])
    encabezado(wd)
    for bloque, explicacion in BLOQUES_DUPLICADAS:
        filas_bloque = [d for d in duplicadas if d["bloque"] == bloque]
        wd.append([f"{bloque} ({len(filas_bloque)} filas): {explicacion}"])
        wd.cell(wd.max_row, 1).font = Font(bold=True)
        for d in filas_bloque:
            g = d["conservada"][0]
            rp, rg = fila_row(d["fila"]), fila_row(g)
            imp_p, imp_g = d["metricas"][0][0]["alcance_o_visualizaciones"], d["metricas"][0][1]["alcance_o_visualizaciones"]
            otras = ", ".join(str(x) for x in d["conservada"])
            if bloque == "Repetición genuina":
                motivo = (f"Es la misma pieza que la fila {otras}: coinciden {CLAVE_TEXTO}, y tiene las mismas métricas. "
                          f"El tablero la muestra una sola vez, con la fila {otras}.")
            elif bloque == "Sin métricas":
                motivo = (f"Coincide con la fila {otras} en {CLAVE_TEXTO}, y ninguna de las dos tiene métricas cargadas. "
                          f"El tablero muestra solo la fila {otras}.")
            else:
                motivo = (f"Coincide con la fila {otras} en {CLAVE_TEXTO}, pero las impresiones son distintas ({_fmt_num(imp_p)} contra {_fmt_num(imp_g)}). "
                          f"El tablero conserva la fila {otras} y no suma las métricas. Confirmar si es la misma pieza cargada dos veces o si son piezas "
                          "distintas; si son distintas, diferenciarlas en la planilla (por ejemplo en el título o el link).")
                if limpio(rp["LINK MATERIAL"]) is None and limpio(rp["TIPO DE MATERIAL"]) is None:
                    motivo += " En las dos filas el LINK MATERIAL y el TIPO DE MATERIAL están vacíos."
            wd.append(datos(d["fila"]) + [bloque, otras, imp_p, imp_g, _por_que_queda(rp, rg), motivo, limpio(rp["COD. CONTENIDO"])])
    wd.freeze_panes = "B2"
    anchos(wd, [10, 26, 40, 34, 30, 40, 12, 22, 14, 16, 16, 34, 80, 16])

    for ws in (wp, wd):
        for r in ws.iter_rows(min_row=2):
            for c in r:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in wd.iter_rows(min_row=2):
        for i in (9, 10):
            r[i].number_format = "#,##0"
    wb.save(path)
    return resumen


def cuadrar_no_procesadas(path, resumen, json_path, stats):
    """Relee el documento y el JSON escritos y se detiene si no coinciden."""
    from openpyxl import load_workbook

    errores = []
    datos = json.load(open(json_path, encoding="utf8"))
    entradas = [p for _, _, p in entradas_de(datos)]
    filas_json = {p["fila_origen"] for p in entradas}
    wb = load_workbook(path, read_only=True)
    res = {}
    for k, v in wb["Resumen"].iter_rows(min_row=2, max_col=2, values_only=True):
        if k and k not in res:
            res[k] = v
    for k, v in resumen:
        if res.get(k) != v:
            errores.append(f"Resumen '{k}': en el documento {res.get(k)!r}, calculado {v!r}")
    esperado = {
        "Entradas del tablero (pautas_barrios.json)": len(entradas),
        "Entran al tablero": len(filas_json),
        "Duplicadas (todas sus entradas se unificaron con otra fila)": stats["filas_reales"] - len(filas_json) - stats["filas_por_destino_excluido"] - stats["filas_por_destino_pendiente"],
    }
    for k, v in esperado.items():
        if res.get(k) != v:
            errores.append(f"Resumen '{k}' = {res.get(k)!r} pero el JSON da {v!r}")

    def filas_hoja(nombre):
        return [r[0] for r in wb[nombre].iter_rows(min_row=2, max_col=1, values_only=True) if isinstance(r[0], int)]

    for hoja, clave in (("Pendientes", "Pendientes"), ("Duplicadas", "Duplicadas (todas sus entradas se unificaron con otra fila)")):
        filas = filas_hoja(hoja)
        if len(filas) != res.get(clave) or len(set(filas)) != len(filas):
            errores.append(f"Hoja {hoja}: {len(filas)} filas ({len(set(filas))} distintas), el Resumen dice {res.get(clave)!r}")
        if set(filas) & filas_json:
            errores.append(f"Hoja {hoja}: filas que también están en el JSON: {sorted(set(filas) & filas_json)[:10]}")
    if errores:
        Path(path).unlink()  # no dejar un documento que dice "OK" sin cuadrar
        raise SystemExit("El documento de filas no procesadas no cuadra con el JSON (no se dejó el documento):\n  " + "\n  ".join(errores))

def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--crudo", default=str(CRUDO_DEFECTO), help="Pauta RRSS_GLOBAL.xlsx (por defecto Contexto/Pauta RRSS_GLOBAL.xlsx)")
    ap.add_argument("--salida", default=str(SALIDA_DEFECTO), help="JSON a escribir (por defecto public/data/pautas_barrios.json)")
    ap.add_argument("--no-procesadas", default=None, help="documento de filas no procesadas (por defecto, junto al JSON)")
    ap.add_argument("--solo-verificar", action="store_true", help="imprime la verificación y no escribe ningún archivo")
    ap.add_argument("--previo", default=None, help="JSON anterior para el freno de bajas (por defecto, el que se va a pisar)")
    ap.add_argument("--aceptar-bajas", action="store_true",
                    help="continuar aunque desaparezcan piezas del JSON anterior (solo si el área confirmó las bajas)")
    ap.add_argument("--hoy", default=date.today().isoformat())
    args = ap.parse_args(argv)

    ref = cargar_referencias()
    df = pd.read_excel(args.crudo, sheet_name="GLOBAL_ImportRange")
    salida, stats, avisos = generar(df, ref)
    piezas = [p for _, _, p in entradas_de(salida)]

    print("--- Verificación Pauta Capital ---")
    print(f"Filas del crudo: {stats['filas_totales']} | descartadas como basura: {stats['filas_descartadas']} | reales: {stats['filas_reales']}")
    print(f"Filas con COD. CONTENIDO (informativo, no decide nada): {stats['filas_con_cod_contenido']} | códigos únicos: {stats['codigos_unicos_con_cod_contenido']}")
    print("Filas por destino:", {k.replace('filas_por_destino_', ''): v for k, v in sorted(stats.items()) if k.startswith('filas_por_destino_')})
    print(f"Entradas antes de dedup: {stats['entradas_antes_de_dedup']} | colapsadas por dedup: {stats['entradas_colapsadas_por_dedup']} | entradas finales: {len(piezas)}")
    print(f"Entradas por bucket: barrios {sum(v['total_pautas'] for v in salida['barrios'].values())} en {len(salida['barrios'])} barrios, "
          f"circuitos_sin_barrio {sum(v['total_pautas'] for v in salida['circuitos_sin_barrio'].values())}, toda_capital {salida['toda_capital']['total_pautas']}")
    conteo_estados = Counter(p["estado"] for p in piezas)
    print("Entradas por estado:", dict(conteo_estados))
    print("Entradas por plataforma:", dict(Counter(p["plataforma"] for p in piezas)))
    no_mapeados = sorted(set(conteo_estados) - ESTADOS_MAPEADOS)
    if no_mapeados:
        print(f"AVISO: estados fuera de ACTIVO/INACTIVO/OBSERVADO/PROGRAMADO/vacío, mostrados tal cual: {no_mapeados}")
    if avisos:
        print(f"AVISO: {len(avisos)} ítems de audiencia sin resolver (la fila entra con el resto):")
        for fila, a in avisos[:20]:
            print(f"   fila {fila}: {a}")
    sin_dato = sum(1 for p in piezas if p["inicio"] is None or p["fin"] is None)
    print(f"Entradas sin INICIO o sin FIN válidos: {sin_dato}")

    # Freno de bajas: una pieza (CAMPAÑA + TITULO DE LA PIEZA, normalizados) que estaba en el JSON
    # anterior y ya no está. Como el crudo es acumulado, una baja es sospechosa. No usa COD. CONTENIDO.
    previo = Path(args.previo or args.salida)
    if previo.exists():
        anterior = json.load(open(previo, encoding="utf8"))
        pieza = lambda p: (normalizar(p.get("campania")), normalizar(p.get("titulo_pieza")))
        bajas = sorted({pieza(p) for _, _, p in entradas_de(anterior)} - {pieza(p) for p in piezas})
        print(f"Piezas (campaña + título) dadas de baja respecto al JSON anterior: {len(bajas)}")
        if bajas:
            print("Detalle:", [" · ".join(b) for b in bajas])
            if args.aceptar_bajas:
                print("--aceptar-bajas: se continúa (las bajas quedan registradas arriba).")
            else:
                raise SystemExit("Aparecieron bajas respecto al JSON anterior. Se detiene la corrida sin escribir: podría ser una pérdida de datos no esperada. Si el área confirmó las bajas, repetir con --aceptar-bajas.")

    if args.solo_verificar:
        print("(--solo-verificar: no se escribió ningún archivo)")
        return
    with open(args.salida, "w", encoding="utf8") as f:
        f.write(json.dumps(salida, ensure_ascii=False, indent=2))
    print(f"Escrito: {args.salida}")
    destino_doc = Path(args.no_procesadas) if args.no_procesadas else Path(args.salida).with_name(NO_PROCESADAS)
    resumen = escribir_no_procesadas(destino_doc, df, stats)
    cuadrar_no_procesadas(destino_doc, resumen, Path(args.salida), stats)
    print(f"Escrito: {destino_doc} (cuadra con el JSON)")
    for k, v in resumen:
        print(f"   {k}: {v}")

if __name__ == "__main__":
    main()
